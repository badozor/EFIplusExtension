"""Excel input template, input reading and output export (manual Part II, 4.2 and 5)."""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import schema as S

HEADER_FILL = PatternFill("solid", fgColor=S.EFI_BLUE_DARK.lstrip("#"))
OPTIONAL_FILL = PatternFill("solid", fgColor="6090C0")
HEADER_FONT = Font(bold=True, color="FFFFFF")


# ---------------------------------------------------------------- template
def build_template(example: bool = False) -> bytes:
    """Blank EFI+ input spreadsheet with drop-down lists (EFI+SpreadSheet)."""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = S.SHEET_INPUT1
    ws2 = wb.create_sheet(S.SHEET_INPUT2)
    lists = wb.create_sheet("lists")

    # lists sheet holds the allowed values used by the drop-downs
    col = 1
    ranges = {}
    for v in S.INPUT_VARS:
        if v.kind == "cat" or v.choices:
            lists.cell(1, col, v.name).font = Font(bold=True)
            for i, c in enumerate(v.choices, start=2):
                lists.cell(i, col, c)
            L = get_column_letter(col)
            ranges[v.name] = f"lists!${L}$2:${L}${len(v.choices) + 1}"
            col += 1
    lists.cell(1, col, "Diadromous").font = Font(bold=True)
    for i, c in enumerate(S.DIAD_CHOICES, start=2):
        lists.cell(i, col, c)
    L = get_column_letter(col)
    ranges["diad"] = f"lists!${L}$2:${L}${len(S.DIAD_CHOICES) + 1}"
    lists.sheet_state = "hidden"

    for j, v in enumerate(S.INPUT_VARS, start=1):
        c = ws1.cell(1, j, v.name)
        c.fill = HEADER_FILL if v.mandatory else OPTIONAL_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        txt = v.label
        if v.unit:
            txt += f" ({v.unit})"
        if v.vmin is not None:
            txt += f"\nRange: {v.vmin:g} to {v.vmax:g}"
        if v.help:
            txt += f"\n{v.help}"
        txt += "\nMissing: " + ("-999" if v.kind in ("int", "float") else "NoData")
        c.comment = Comment(txt, "EFI+")
        ws1.column_dimensions[get_column_letter(j)].width = max(12, min(28, len(v.name) + 4))
        L = get_column_letter(j)
        if v.name in ranges:
            dv = DataValidation(type="list", formula1=f"={ranges[v.name]}", allow_blank=True,
                                showErrorMessage=v.name != "Species")
            dv.error = f"Allowed values for {v.name}: see the drop-down list."
            ws1.add_data_validation(dv)
            dv.add(f"{L}2:{L}5000")
        elif v.kind in ("int", "float") and v.vmin is not None:
            dv = DataValidation(type="decimal" if v.kind == "float" else "whole", operator="between",
                                formula1=str(v.vmin), formula2=str(v.vmax), allow_blank=True)
            dv.error = f"{v.name}: {v.vmin:g} to {v.vmax:g} (or -999 if missing)."
            dv.errorStyle = "warning"
            ws1.add_data_validation(dv)
            dv.add(f"{L}2:{L}5000")
    ws1.freeze_panes = "A2"
    ws1.row_dimensions[1].height = 30

    for j, name in enumerate(S.DIAD_KEY_COLUMNS + S.DIAD_SPECIES_COLUMNS, start=1):
        c = ws2.cell(1, j, name)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws2.column_dimensions[get_column_letter(j)].width = 14
        if name in S.DIAD_SPECIES_COLUMNS:
            L = get_column_letter(j)
            dv = DataValidation(type="list", formula1=f"={ranges['diad']}", allow_blank=True)
            ws2.add_data_validation(dv)
            dv.add(f"{L}2:{L}5000")
    ws2.freeze_panes = "G2"
    ws2.row_dimensions[1].height = 45

    if example:
        from .examples import example_input1, example_input2
        for i, row in enumerate(example_input1().itertuples(index=False), start=2):
            for j, val in enumerate(row, start=1):
                ws1.cell(i, j, val)
        e2 = example_input2()[S.DIAD_KEY_COLUMNS + S.DIAD_SPECIES_COLUMNS]
        for i, row in enumerate(e2.itertuples(index=False), start=2):
            for j, val in enumerate(row, start=1):
                ws2.cell(i, j, val)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- reading
def read_workbook(data: bytes | str) -> tuple[pd.DataFrame, pd.DataFrame | None, list[str]]:
    """Read input1 (and input2 if present) from an .xlsx/.xls file or a ;-separated text file."""
    notes: list[str] = []
    if isinstance(data, str):
        with open(data, "rb") as f:
            data = f.read()
    head = data[:8]
    if head.startswith(b"PK") or head.startswith(b"\xd0\xcf\x11\xe0"):
        engine = None if head.startswith(b"PK") else "xlrd"
        try:
            sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype=object, engine=engine)
        except ImportError:
            raise ValueError("Old .xls files need the 'xlrd' package. Save the file as .xlsx or install xlrd.")
        names = list(sheets)
        low = {n.lower().strip(): n for n in names}
        s1 = low.get(S.SHEET_INPUT1.lower()) or next((n for n in names if n.lower().startswith("input1")), names[0])
        s2 = low.get(S.SHEET_INPUT2.lower()) or next((n for n in names if n.lower().startswith("input2")), None)
        if s2 is None and len(names) > 1 and names[1] != "lists":
            s2 = names[1]
        notes.append(f"Sheet used for input1: '{s1}'.")
        df2 = sheets[s2] if s2 else None
        if s2:
            notes.append(f"Sheet used for input2: '{s2}'.")
        if df2 is not None and df2.dropna(how="all").empty:
            df2 = None
        return sheets[s1], df2, notes
    # text file (Input.txt as produced by the web application)
    text = data.decode("utf-8", errors="replace")
    df = pd.read_csv(io.StringIO(text), sep=";", dtype=object)
    notes.append("Text file read as input1 (semicolon separated).")
    return df, None, notes


def frames_to_excel_copy(input1: pd.DataFrame, input2: pd.DataFrame | None) -> bytes:
    """Excel copy of the input data (manual Figure 10), re-importable in the GUI."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        input1.to_excel(xw, sheet_name=S.SHEET_INPUT1, index=False)
        if input2 is not None and not input2.empty:
            input2[S.DIAD_KEY_COLUMNS + S.DIAD_SPECIES_COLUMNS].to_excel(xw, sheet_name=S.SHEET_INPUT2, index=False)
    return buf.getvalue()


# ---------------------------------------------------------------- output
def order_output(df: pd.DataFrame) -> pd.DataFrame:
    cols = [c for g in S.OUTPUT_GROUPS.values() for c in g if c in df.columns]
    rest = [c for c in df.columns if c not in cols]
    return df[cols + rest]


def build_output_workbook(output: pd.DataFrame, input1: pd.DataFrame, input2: pd.DataFrame | None,
                          info: str, meta: dict) -> bytes:
    """NewEFI+output workbook: grouped output table, inputs and run information."""
    out = order_output(output)
    wb = Workbook()
    ws = wb.active
    ws.title = "EFI+ output"
    group_fills = ["DDEBF7", "E2EFDA", "FFF2CC", "FCE4D6", "EDEDED"]
    # row 1: group labels, row 2: variable names
    j = 1
    for (g, cols), color in zip(S.OUTPUT_GROUPS.items(), group_fills):
        present = [c for c in cols if c in out.columns]
        if not present:
            continue
        ws.cell(1, j, g).font = Font(bold=True)
        ws.merge_cells(start_row=1, start_column=j, end_row=1, end_column=j + len(present) - 1)
        for k in range(len(present)):
            ws.cell(1, j + k).fill = PatternFill("solid", fgColor=color)
            ws.cell(2, j + k).fill = PatternFill("solid", fgColor=color)
        j += len(present)
    for j, c in enumerate(out.columns, start=1):
        cell = ws.cell(2, j, c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(j)].width = 14 if not c.startswith("Comments") else 26
    cls_col = list(out.columns).index("FishIndex.class") + 1 if "FishIndex.class" in out.columns else None
    for i, row in enumerate(out.itertuples(index=False), start=3):
        for j, val in enumerate(row, start=1):
            if isinstance(val, float) and pd.isna(val):
                val = None
            elif isinstance(val, float):
                val = round(val, 4)
            ws.cell(i, j, val)
        if cls_col:
            v = ws.cell(i, cls_col).value
            try:
                k = int(v)
                ws.cell(i, cls_col).fill = PatternFill("solid", fgColor=S.CLASS_COLORS[k].lstrip("#"))
                ws.cell(i, cls_col).font = Font(bold=True, color="FFFFFF" if k in (1, 2, 5) else "000000")
            except (TypeError, ValueError, KeyError):
                pass
    ws.freeze_panes = "B3"
    ws.row_dimensions[2].height = 45

    leg = wb.create_sheet("Class legend")
    leg.append(["FishIndex.class", "Ecological status"])
    for k, lab in S.CLASS_LABELS.items():
        leg.append([k, lab])
        leg.cell(leg.max_row, 1).fill = PatternFill("solid", fgColor=S.CLASS_COLORS[k].lstrip("#"))

    for name, df in (("input1", input1), ("input2", input2)):
        if df is not None and not df.empty:
            w = wb.create_sheet(name)
            w.append(list(df.columns))
            for row in df.itertuples(index=False):
                w.append(list(row))

    w = wb.create_sheet("Run information")
    w.append(["EFI+ computation"])
    w["A1"].font = Font(bold=True)
    for k, v in meta.items():
        w.append([k, str(v)])
    w.append([])
    w.append(["R log (info.out)"])
    for line in info.splitlines():
        w.append([line])
    w.column_dimensions["A"].width = 40
    w.column_dimensions["B"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")
